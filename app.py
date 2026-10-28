
import streamlit as st
import pandas as pd
import fitz
import plotly.express as px
from docx import Document
from io import BytesIO
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity


# ==================================================
# RESUMERANK
# ==================================================
st.set_page_config(
    page_title="ResumeRank",
    page_icon="📄",
    layout="wide"
)

if "results" not in st.session_state:
    st.session_state.results = None


# ==================================================
# LIGHT CREAM THEME
# ==================================================
BG = "#FFF9EF"
PANEL = "#FFF1DF"
TEXT = "#352F27"
MUTED = "#786F63"
BORDER = "#E5D7C3"
FIELD = "#FFFCF6"
UPLOAD_BG = "#FFF4E5"
ACCENT = "#805A32"

# Chart colors only
CHART_BG = "#F7F2E8"
CHART_BAR = "#76543D"
CHART_GRID = "#E0D5C7"


# ==================================================
# CUSTOM CSS - ORIGINAL DESIGN
# ==================================================
st.markdown(
    f"""
    <style>
    /* Main page */
    .stApp,
    [data-testid="stAppViewContainer"] {{
        background-color: {BG} !important;
        color: {TEXT} !important;
    }}

    [data-testid="stHeader"] {{
        background-color: {BG} !important;
    }}

    [data-testid="stSidebar"] {{
        background-color: {PANEL} !important;
    }}

    /* Text */
    .stApp h1,
    .stApp h2,
    .stApp h3,
    .stApp h4,
    .stApp p,
    .stApp label,
    .stApp li,
    .stApp [data-testid="stCaptionContainer"] p {{
        color: {TEXT} !important;
    }}

    /* Tabs */
    .stApp [data-testid="stTabs"] button[role="tab"],
    .stApp [data-testid="stTabs"] button[role="tab"] p,
    .stApp [data-testid="stTabs"] button[role="tab"] span {{
        color: {TEXT} !important;
        opacity: 1 !important;
    }}

    .stApp [data-testid="stTabs"] button[role="tab"][aria-selected="true"],
    .stApp [data-testid="stTabs"] button[role="tab"][aria-selected="true"] p,
    .stApp [data-testid="stTabs"] button[role="tab"][aria-selected="true"] span {{
        color: {ACCENT} !important;
    }}

    /* Job description */
    .stApp textarea,
    .stApp [data-baseweb="textarea"] textarea {{
        background-color: {FIELD} !important;
        color: {TEXT} !important;
        -webkit-text-fill-color: {TEXT} !important;
        border: 1px solid {BORDER} !important;
        border-radius: 8px !important;
        caret-color: {TEXT} !important;
    }}

    .stApp textarea::placeholder {{
        color: {MUTED} !important;
        -webkit-text-fill-color: {MUTED} !important;
    }}

    /* Upload area */
    .stApp [data-testid="stFileUploaderDropzone"] {{
        background-color: {UPLOAD_BG} !important;
        border: 1px dashed {BORDER} !important;
        border-radius: 12px !important;
    }}

    .stApp [data-testid="stFileUploaderDropzone"] * {{
        color: {TEXT} !important;
    }}

    .stApp [data-testid="stFileUploaderDropzone"] button {{
        background-color: {PANEL} !important;
        color: {TEXT} !important;
        border: 1px solid {BORDER} !important;
        border-radius: 7px !important;
    }}

    /* Uploaded files */
    .stApp [data-testid="stFileUploader"] [data-testid="stFile"],
    .stApp [data-testid="stFileUploader"] [data-testid="stFileUploaderFile"] {{
        background-color: {PANEL} !important;
        color: {TEXT} !important;
        border: 1px solid {BORDER} !important;
    }}

    .stApp [data-testid="stFileUploader"] [data-testid="stFile"] *,
    .stApp [data-testid="stFileUploader"] [data-testid="stFileUploaderFile"] * {{
        color: {TEXT} !important;
    }}

    /* Primary action button */
    .stApp button[kind="primary"] {{
        background-color: {ACCENT} !important;
        color: #FFFFFF !important;
        border: 1px solid {ACCENT} !important;
        border-radius: 8px !important;
        min-height: 42px;
    }}

    .stApp button[kind="primary"] * {{
        color: #FFFFFF !important;
        -webkit-text-fill-color: #FFFFFF !important;
    }}

    /* Other and download buttons */
    .stApp button[kind="secondary"],
    .stApp .stDownloadButton button {{
        background-color: {PANEL} !important;
        color: {TEXT} !important;
        border: 1px solid {BORDER} !important;
        border-radius: 8px !important;
    }}

    .stApp button[kind="secondary"] *,
    .stApp .stDownloadButton button * {{
        color: {TEXT} !important;
    }}

    /* Metrics */
    .stApp [data-testid="stMetric"] {{
        background-color: {PANEL} !important;
        border: 1px solid {BORDER} !important;
        padding: 14px;
        border-radius: 10px;
    }}

    .stApp [data-testid="stMetricLabel"],
    .stApp [data-testid="stMetricValue"],
    .stApp [data-testid="stMetricDelta"] {{
        color: {TEXT} !important;
    }}

    /* Tables */
    .stApp [data-testid="stDataFrame"],
    .stApp [data-testid="stTable"] {{
        border: 1px solid {BORDER};
        border-radius: 8px;
    }}

    .stApp hr {{
        border-color: {BORDER} !important;
    }}

    /* Author footer */
    .author-footer {{
        text-align: center;
        color: {MUTED};
        font-size: 13px;
        padding: 20px 0 15px;
        border-top: 1px solid {BORDER};
        margin-top: 30px;
    }}
    </style>
    """,
    unsafe_allow_html=True
)


# ==================================================
# HEADER
# ==================================================
st.title("ResumeRank")
st.caption("AI-powered Resume Screening and Candidate Ranking")


# ==================================================
# TEXT EXTRACTION
# ==================================================
def extract_pdf_text(file_bytes):
    with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
        return "\n".join(page.get_text() for page in pdf).strip()


def extract_docx_text(file_bytes):
    document = Document(BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))

    return "\n".join(parts).strip()


def extract_resume_text(uploaded_file):
    data = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        return extract_pdf_text(data)

    if filename.endswith(".docx"):
        return extract_docx_text(data)

    return ""


# ==================================================
# AI MODEL
# ==================================================
@st.cache_resource
def load_model():
    return SentenceTransformer(
        "sentence-transformers/bert-base-nli-mean-tokens"
    )


# ==================================================
# CHART STYLE - ONLY FOR HR ANALYSIS CHARTS
# ==================================================
def style_chart(fig):
    fig.update_layout(
        paper_bgcolor=CHART_BG,
        plot_bgcolor=CHART_BG,
        font=dict(
            family="Arial",
            color=TEXT,
            size=12
        ),
        margin=dict(l=20, r=20, t=35, b=65),
        showlegend=False,
        xaxis=dict(
            title_font=dict(color=TEXT),
            tickfont=dict(color=TEXT),
            gridcolor=CHART_GRID,
            linecolor=CHART_GRID,
            zerolinecolor=CHART_GRID
        ),
        yaxis=dict(
            title_font=dict(color=TEXT),
            tickfont=dict(color=TEXT),
            gridcolor=CHART_GRID,
            linecolor=CHART_GRID,
            zerolinecolor=CHART_GRID
        )
    )
    return fig


# ==================================================
# MAIN TABS
# ==================================================
screening_tab, analytics_tab = st.tabs(
    ["📄 Resume Screening", "📊 HR Analysis & Download"]
)


# ==================================================
# TAB 1: RESUME SCREENING - ORIGINAL
# ==================================================
with screening_tab:
    st.subheader("Screen Candidates")

    job_description = st.text_area(
        "Job Description",
        placeholder="Paste the job description here...",
        height=180
    )

    uploaded_files = st.file_uploader(
        "Upload Candidate Resumes",
        type=["pdf", "docx"],
        accept_multiple_files=True,
        help="Upload PDF or DOCX resumes."
    )

    if st.button(
        "Screen Resumes",
        type="primary",
        use_container_width=True
    ):
        if not job_description.strip():
            st.warning("Please enter a job description.")

        elif not uploaded_files:
            st.warning("Please upload at least one resume.")

        else:
            names = []
            resume_texts = []
            errors = []

            with st.spinner("Reading and analysing resumes..."):
                for file in uploaded_files:
                    try:
                        content = extract_resume_text(file)

                        if content:
                            names.append(file.name)
                            resume_texts.append(content)
                        else:
                            errors.append(
                                f"{file.name}: No readable text found."
                            )

                    except Exception:
                        errors.append(
                            f"{file.name}: Unable to read this file."
                        )

                for error in errors:
                    st.warning(error)

                if not resume_texts:
                    st.error(
                        "No readable resumes found. Please check your files."
                    )

                else:
                    try:
                        model = load_model()

                        job_vector = model.encode(
                            [job_description],
                            convert_to_numpy=True
                        )

                        resume_vectors = model.encode(
                            resume_texts,
                            convert_to_numpy=True,
                            show_progress_bar=False
                        )

                        scores = cosine_similarity(
                            resume_vectors,
                            job_vector
                        ).flatten()

                        results = pd.DataFrame({
                            "Candidate": names,
                            "Match Score (%)": (scores * 100).round(2)
                        })

                        results = results.sort_values(
                            "Match Score (%)",
                            ascending=False
                        ).reset_index(drop=True)

                        results.insert(
                            0, "Rank", range(1, len(results) + 1)
                        )

                        results["Status"] = results[
                            "Match Score (%)"
                        ].apply(
                            lambda score:
                            "Shortlisted" if score >= 50 else "Review"
                        )

                        st.session_state.results = results
                        st.success("Resume screening completed!")

                    except Exception:
                        st.error(
                            "The AI model could not process the resumes. "
                            "Check your internet connection and try again."
                        )

    if st.session_state.results is not None:
        st.divider()
        st.subheader("Candidate Rankings")

        st.dataframe(
            st.session_state.results,
            use_container_width=True,
            hide_index=True
        )

        st.caption(
            "Scores indicate text similarity, not guaranteed job suitability. "
            "Review resumes before making hiring decisions."
        )


# ==================================================
# TAB 2: HR ANALYSIS & DOWNLOAD
# Only chart rendering is changed
# ==================================================
with analytics_tab:
    st.subheader("HR Analysis")

    if st.session_state.results is None:
        st.info(
            "Screen resumes first to view analytics and download results."
        )

    else:
        results = st.session_state.results

        total = len(results)
        shortlisted = int((results["Status"] == "Shortlisted").sum())
        review = total - shortlisted
        average = results["Match Score (%)"].mean()

        col1, col2, col3, col4 = st.columns(4)
        col1.metric("Total Candidates", total)
        col2.metric("Shortlisted", shortlisted)
        col3.metric("Needs Review", review)
        col4.metric("Average Match", f"{average:.1f}%")

        st.divider()

        chart1, chart2 = st.columns(2)

        # Chart 1: Candidate Match Scores
        with chart1:
            st.subheader("Candidate Match Scores")

            score_data = results.sort_values(
                "Match Score (%)",
                ascending=False
            )

            fig_scores = px.bar(
                score_data,
                x="Candidate",
                y="Match Score (%)",
                color_discrete_sequence=[CHART_BAR],
                text="Match Score (%)"
            )

            fig_scores.update_traces(
                marker_color=CHART_BAR,
                texttemplate="%{text:.1f}%",
                textposition="outside",
                cliponaxis=False
            )

            fig_scores.update_yaxes(
                title_text="Match Score (%)",
                rangemode="tozero"
            )

            fig_scores.update_xaxes(
                title_text="Candidate",
                tickangle=-35
            )

            fig_scores = style_chart(fig_scores)

            st.plotly_chart(
                fig_scores,
                use_container_width=True,
                config={"displayModeBar": False}
            )

        # Chart 2: Candidate Status
        with chart2:
            st.subheader("Candidate Status")

            counts = results["Status"].value_counts().reindex(
                ["Shortlisted", "Review"],
                fill_value=0
            )

            status_data = pd.DataFrame({
                "Status": counts.index,
                "Candidates": counts.values
            })

            fig_status = px.bar(
                status_data,
                x="Status",
                y="Candidates",
                color_discrete_sequence=[CHART_BAR],
                text="Candidates"
            )

            fig_status.update_traces(
                marker_color=CHART_BAR,
                textposition="outside",
                cliponaxis=False
            )

            fig_status.update_yaxes(
                title_text="Number of Candidates",
                rangemode="tozero",
                dtick=1
            )

            fig_status.update_xaxes(
                title_text="Candidate Status"
            )

            fig_status = style_chart(fig_status)

            st.plotly_chart(
                fig_status,
                use_container_width=True,
                config={"displayModeBar": False}
            )

        st.divider()
        st.subheader("Download Results")

        csv_data = results.to_csv(index=False).encode("utf-8-sig")

        st.download_button(
            "⬇️ Download Results as CSV",
            data=csv_data,
            file_name="ResumeRank_Candidate_Results.csv",
            mime="text/csv",
            use_container_width=True
        )

        st.caption(
            "Review the original resumes and evaluate candidates fairly."
        )
